"""La historia de usuario en Word, con el formato de la plantilla: banner del Feature, id y título, Como/Quiero/Para y criterios.

No incluye citas ni datos de las reuniones: eso se consulta en la pantalla. Usa python-docx (dependencia opcional «export»).
"""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from backlogguard.export.story import Story

BANNER, US_ID, LABEL, ROW, CRITERIA_HEADER = "4B4BA0", "5B6B7B", "B4B7DC", "DDE6F5", "7DB383"
DARK, WHITE, GREEN, BORDER = "1F2A6B", "FFFFFF", "2E9E5B", "9FA8C4"
FONT = "Arial"
WIDTH = 17.0  # cm, A4 con márgenes de 2


def _shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def _borders(table) -> None:
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), BORDER)
        borders.append(element)
    props.append(borders)


def _write(cell, runs: list[tuple[str, bool, str]], *, size: float = 11, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    """Escribe en una celda uno o varios tramos (texto, negrita, color)."""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    for text, bold, color in runs:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)


def _table(document, rows: int, cols: int, widths: list[float]):
    table = document.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
    return table


def story_to_docx(story: Story) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    document.styles["Normal"].font.name = FONT

    banner = _table(document, 1, 1, [WIDTH])
    _shade(banner.cell(0, 0), BANNER)
    label = f"Feature {story.feature_number}: {story.feature}" if story.feature and story.feature_number else (story.feature or "Historia de usuario")
    _write(banner.cell(0, 0), [(label, True, WHITE)], size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

    rows = 1 + (1 if story.feature else 0) + 3 + 1 + len(story.criteria)
    table = _table(document, rows, 2, [3.2, WIDTH - 3.2])
    at = 0
    _shade(table.cell(at, 0), US_ID)
    _shade(table.cell(at, 1), BANNER)
    _write(table.cell(at, 0), [(story.us_id, True, WHITE)], size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    _write(table.cell(at, 1), [(story.title, True, WHITE)], size=13)
    if story.feature:
        at += 1
        merged = table.cell(at, 0).merge(table.cell(at, 1))
        _shade(merged, ROW)
        _write(merged, [("Feature: ", True, DARK), (story.feature, False, "000000")], size=12)
    for name, text in (("Como...", story.como), ("Quiero...", story.quiero), ("Para...", story.para)):
        at += 1
        _shade(table.cell(at, 0), LABEL)
        _shade(table.cell(at, 1), ROW)
        _write(table.cell(at, 0), [(name, True, WHITE)], size=12)
        _write(table.cell(at, 1), [(text, False, "000000")], size=12)
    at += 1
    header = table.cell(at, 0).merge(table.cell(at, 1))
    _shade(header, CRITERIA_HEADER)
    _write(header, [("Criterios de Aceptación", True, WHITE)], size=12)
    for criterion in story.criteria:
        at += 1
        _shade(table.cell(at, 0), ROW)
        _shade(table.cell(at, 1), ROW)
        _write(table.cell(at, 0), [("✔", True, GREEN)], size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(table.cell(at, 1), [(criterion, False, "000000")], size=11)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
