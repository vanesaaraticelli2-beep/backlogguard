"""La historia de usuario completa: Como/Quiero/Para, escenarios y reglas, Feature y el Word por historia."""
from io import BytesIO

import pytest
from docx import Document
from test_catalogo import approve_and_apply, draft, evidence
from test_continuidad import world  # noqa: F401  (fixture con catálogo y línea base)
from test_graph import EMAIL, GHERKIN, STORY, World, tool

from backlogguard.board import BacklogItem, Board, Citation, ItemKind, SourceType
from backlogguard.catalog import Proposal, ProposalKind
from backlogguard.export import format_criterion, split_story, story_of, story_to_docx
from backlogguard.export.story import us_id
from backlogguard.graph.drafting import GOAL, context_of
from backlogguard.tools.requerimientos import DraftArgs, RequirementsToolkit
from backlogguard.ui.story_view import story_html

QUOTE = "Esta frase solo vive en la reunión y no va al Word."
SCENARIO = "Dado un asesor asignado Cuando intento asignarlo a otra alianza Entonces el sistema rechaza la asignación"
RULE = "El campo acepta solo caracteres numéricos"


def item(**extra):
    base = dict(
        id="HU-RQ-003", kind=ItemKind.HISTORIA, title="Un asesor solo puede pertenecer a una alianza",
        como="administrador del portal", quiero="que el sistema no permita asignar un asesor a más de una alianza", para="evitar vínculos duplicados",
        acceptance_criteria=[SCENARIO, RULE], requirement_id="RQ-003", requirement_version=1,
        citations=[Citation(fragment_id="r:1", quote=QUOTE, source_name="Reunión de inicio", source_type=SourceType.REUNION, source_date="2026-08-24", speaker="Ana")],
    )
    return BacklogItem(**{**base, **extra})


def test_el_como_quiero_para_se_separa_y_los_escenarios_se_leen_en_una_frase():
    assert split_story("Como cliente quiero ver mi pedido para saber cuándo llega.") == ("cliente", "ver mi pedido", "saber cuándo llega")
    assert split_story("Quiero entrar.") is None
    assert format_criterion(SCENARIO) == "Dado un asesor asignado, cuando intento asignarlo a otra alianza, entonces el sistema rechaza la asignación."
    assert format_criterion(RULE) == "El campo acepta solo caracteres numéricos."


def test_el_id_us_sale_del_requisito_y_la_historia_vieja_se_lee_de_su_descripcion():
    assert us_id(item()) == "US-003" and us_id(item(requirement_id=None)) == "HU-RQ-003"
    old = item(como=None, quiero=None, para=None, description="Como administrador quiero bloquear altas para evitar errores.")
    story = story_of(old)
    assert (story.como, story.quiero, story.para) == ("administrador", "bloquear altas", "evitar errores")


def test_el_feature_sale_del_area_de_la_linea_base_del_requisito(world):  # noqa: F811
    catalog, baseline = world["catalog"], world["baseline"]
    proposal = Proposal(
        source_id="reunion_01", fingerprint="fp", candidate_key="k", kind=ProposalKind.NUEVO, drafts=[draft("Vigencia del score", commitment_ids=["C-002"])],
        evidence=[evidence(source_id="reunion_01", fp="fp")],
    )
    rid = approve_and_apply(catalog, catalog.propose(proposal)).affected[0]
    story = story_of(item(requirement_id=rid), catalog, baseline)
    assert (story.feature, story.feature_number) == ("Integración NOSIS", 2)
    assert story_of(item(requirement_id=rid)).feature is None  # sin línea base no se inventa


def test_el_word_tiene_el_formato_de_la_plantilla_y_no_lleva_citas_ni_datos_de_la_reunion():
    story = story_of(item())
    story = type(story)(**{**story.__dict__, "feature": "Administración", "feature_number": 9})
    document = Document(BytesIO(story_to_docx(story)))
    text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    for expected in ("Feature 9: Administración", "US-003", "Un asesor solo puede pertenecer a una alianza", "Como...", "Quiero...", "Para...",
                     "administrador del portal", "Criterios de Aceptación", "Dado un asesor asignado, cuando intento asignarlo", "El campo acepta solo caracteres numéricos.", "✔"):
        assert expected in text, expected
    for absent in ("Canal", "Ramos", QUOTE, "Reunión de inicio", "Ana", "2026-08-24", "Trazabilidad", "RQ-003"):
        assert absent not in text and absent not in "\n".join(p.text for p in document.paragraphs), absent
    assert story.file_name == "US-003 - Un asesor solo puede pertenecer a una alianza.docx"


def test_la_pantalla_escapa_el_texto_y_muestra_todos_los_criterios():
    html = story_html(story_of(item(title="Alta <script>alert(1)</script>")))
    assert "<script>" not in html and "&lt;script&gt;" in html and html.count("us-row us-crit") == 2 and "Criterios de Aceptación" in html


@pytest.fixture
def toolkit(tmp_path):
    w = World(tmp_path, supervisor=[], producer=[], critic=[], auditor=[])
    board = Board()
    return RequirementsToolkit(w.deps.library, w.retriever, board), board, w


def draft_args(w, **extra):
    chunk = w.chunk(*EMAIL)
    return {"kind": "historia", "title": "Notificación por email del estado", "description": "", "citations": [{"chunk_id": chunk.chunk_id, "quote": EMAIL[1]}], **extra}


def test_una_historia_necesita_un_escenario_pero_acepta_reglas_y_arma_su_descripcion_con_como_quiero_para(toolkit):
    tk, board, w = toolkit
    with pytest.raises(ValueError, match="al menos un escenario"):
        DraftArgs(**draft_args(w, acceptance_criteria=[RULE]))
    with pytest.raises(ValueError, match="juntos"):
        DraftArgs(**draft_args(w, como="cliente", acceptance_criteria=[GHERKIN]))
    args = DraftArgs(**draft_args(w, como="cliente", quiero="recibir un email", para="enterarme del cambio", acceptance_criteria=[GHERKIN, RULE, RULE + " también"])).model_dump(mode="json")
    result = tk.redactar_historia(args)
    assert result.ok
    saved = board.items[result.metadata["item_id"]]
    assert saved.description == "Como cliente quiero recibir un email para enterarme del cambio." and (saved.como, saved.quiero, saved.para) == ("cliente", "recibir un email", "enterarme del cambio")
    assert len(saved.acceptance_criteria) == 3


def test_el_productor_recibe_el_contexto_de_la_reunion_y_la_orden_de_cubrir_todos_los_escenarios(world, tmp_path):  # noqa: F811
    assert "TODOS los escenarios" in GOAL and "como" in GOAL and "Dado/Cuando/Entonces" in GOAL
    from backlogguard.graph.deps import index_library
    from backlogguard.rag import DeterministicLexicalEmbeddings, InMemoryRetriever

    retriever = InMemoryRetriever(DeterministicLexicalEmbeddings())
    index_library(world["library"], retriever)
    chunk = next(c for c in retriever.all_chunks({"document_id": "reunion_01"}))
    catalog = world["catalog"]
    proposal = Proposal(
        source_id="reunion_01", fingerprint="fp", candidate_key="k", kind=ProposalKind.NUEVO, drafts=[draft("Validar el CUIT")],
        evidence=[evidence(fragment=chunk.chunk_id, source_id="reunion_01", fp="fp")],
    )
    rid = approve_and_apply(catalog, catalog.propose(proposal)).affected[0]
    context = context_of(catalog.requirement(rid), retriever)
    assert "CUIT" in context and "score de NOSIS" in context  # el fragmento citado y su vecino
    assert context_of(catalog.requirement(rid), None) == ""
